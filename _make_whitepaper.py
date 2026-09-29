from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUT = Path(r'D:\王飞飞\codex项目集\速报系统1.0\output')
OUT.mkdir(exist_ok=True)
DOCX = OUT / '点点速报系统商业版产品白皮书.docx'

img1 = Path(r'D:\谷歌下载\速报系统产品相关架构图-第 1 页.drawio.png')
img2 = Path(r'D:\谷歌下载\速报系统产品相关架构图-第 2 页.drawio.png')
img3 = Path(r'D:\谷歌下载\速报系统功能架构图.png')

BLUE = '1D61D6'
LIGHT_BLUE = 'EAF2FF'
GRAY = '6B7280'
DARK = '111827'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('left', 'top', 'right', 'bottom'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:%s' % edge
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ['val', 'sz', 'space', 'color']:
                if key in edge_data:
                    element.set(qn('w:%s' % key), str(edge_data[key]))


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_run_font(run, size=11, bold=False, color=None, name='Microsoft YaHei'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bullet=False, bold=False, color=DARK, size=11, indent=0, after=4):
    p = doc.add_paragraph(style='List Bullet' if bullet else None)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    set_paragraph_spacing(p, before=0, after=after, line=1.35)
    if bullet:
        p.style = doc.styles['List Bullet']
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=6, line=1.1)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color=BLUE)
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=BLUE)
    else:
        set_run_font(run, size=12, bold=True, color=BLUE)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=2, after=8, line=1.0)
    run = p.add_run(text)
    set_run_font(run, size=10, bold=True, color=GRAY)
    return p


def add_picture_page(doc, path, caption, width_inches=6.4):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=4, line=1.0)
    try:
        p.add_run().add_picture(str(path), width=Inches(width_inches))
    except Exception as e:
        t = doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = t.add_run(f'图片加载失败：{path.name}\n{e}')
        set_run_font(run, size=10, color='C00000')
    add_caption(doc, caption)


def add_kv_table(doc, rows, widths=(3.0, 12.2)):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for left, right in rows:
        row = table.add_row()
        row.cells[0].width = Cm(widths[0])
        row.cells[1].width = Cm(widths[1])
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].text = left
        row.cells[1].text = right
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_shading(row.cells[1], 'FFFFFF')
        for cell in row.cells:
            set_cell_border(cell, left={"val":"single","sz":8,"color":"C9D8F2"}, right={"val":"single","sz":8,"color":"C9D8F2"}, top={"val":"single","sz":8,"color":"C9D8F2"}, bottom={"val":"single","sz":8,"color":"C9D8F2"})
            for p in cell.paragraphs:
                set_paragraph_spacing(p, before=0, after=2, line=1.2)
                for r in p.runs:
                    set_run_font(r, size=10, bold=False, color=DARK)
    return table


def add_feature_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ['终端', '典型角色', '核心能力', '商业定位']
    widths = [2.8, 3.2, 7.0, 3.4]
    for i, w in enumerate(widths):
        table.rows[0].cells[i].width = Cm(w)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, left={"val":"single","sz":10,"color":"B7C7EA"}, right={"val":"single","sz":10,"color":"B7C7EA"}, top={"val":"single","sz":10,"color":"B7C7EA"}, bottom={"val":"single","sz":"10","color":"B7C7EA"})
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, size=10.5, bold=True, color='FFFFFF')
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            cell.text = text
            set_cell_border(cell, left={"val":"single","sz":8,"color":"D5DEF0"}, right={"val":"single","sz":8,"color":"D5DEF0"}, top={"val":"single","sz":8,"color":"D5DEF0"}, bottom={"val":"single","sz":8,"color":"D5DEF0"})
            for p in cell.paragraphs:
                set_paragraph_spacing(p, before=0, after=2, line=1.2)
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    set_run_font(r, size=9.5, bold=False, color=DARK)
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = Cm(w)
    return table


def set_doc_font_styles(doc):
    styles = doc.styles
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        if name in styles:
            st = styles[name]
            st.font.name = 'Microsoft YaHei'
            st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    styles['Normal'].font.size = Pt(10.5)
    styles['Title'].font.size = Pt(22)
    styles['Title'].font.bold = True
    styles['Title'].font.color.rgb = RGBColor.from_string(BLUE)
    styles['Heading 1'].font.size = Pt(16)
    styles['Heading 1'].font.bold = True
    styles['Heading 1'].font.color.rgb = RGBColor.from_string(BLUE)
    styles['Heading 2'].font.size = Pt(13)
    styles['Heading 2'].font.bold = True
    styles['Heading 2'].font.color.rgb = RGBColor.from_string(BLUE)


doc = Document()
set_doc_font_styles(doc)

# Page setup
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

# Footer
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run('点点速报系统商业版产品白皮书  |  内部商业资料')
set_run_font(fr, size=9, color=GRAY)

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=140, after=18, line=1.0)
r = p.add_run('点点速报系统商业版产品白皮书')
set_run_font(r, size=24, bold=True, color=BLUE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=2, after=20, line=1.0)
r = p.add_run('统一入口门户 · H5报送审核端 · V8客户管理端 · MT应用管理端')
set_run_font(r, size=12, bold=True, color=DARK)

rows = [
    ('产品定位', '面向多级组织的速报上报、AI辅助审核、统计考核与运营管理一体化平台'),
    ('交付形态', '统一入口门户 + 三大业务端 + 可扩展的能力与数据底座'),
    ('部署模式', '支持私有化部署、专属云部署与混合部署'),
    ('适用客户', '政企组织、园区平台、集团型机构、需要分级审核与留痕管理的业务场景'),
    ('编制版本', '商业版 / V1.0'),
    ('编制日期', '2026-08-25'),
]
add_kv_table(doc, rows)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=16, after=6, line=1.0)
r = p.add_run('本白皮书基于现有产品架构图、业务流程图、功能架构图及已实现的三端项目整理而成。')
set_run_font(r, size=10.5, color=GRAY)

# Contents page
doc.add_page_break()
add_heading(doc, '目录', 1)
content_items = [
    '一、概述',
    '二、行业现状',
    '三、产品介绍',
    '  3.1 Web客户端',
    '  3.2 H5端',
    '  3.3 人工服务端',
    '四、核心功能（市场商务方向）',
    '  4.1 统一入口与多端协同工作台',
    '  4.2 AI驱动的速报生成与审核闭环',
    '  4.3 多级组织、租户与权限运营',
    '  4.4 业务流程与状态流转',
    '五、适用场景',
    '六、部署与服务模式',
    '七、产品优势',
    '八、客户价值',
    '九、结语',
    '附录A 产品总体架构图',
    '附录B 业务流程与状态流转图',
    '附录C 功能架构图',
]
for item in content_items:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4 if not item.startswith('  ') else 1.0)
    set_paragraph_spacing(p, before=0, after=2, line=1.15)
    run = p.add_run(item)
    set_run_font(run, size=11 if not item.startswith('  ') else 10.5, color=DARK)

# Section 1
add_heading(doc, '一、概述', 1)
add_body(doc, '点点速报系统是面向多级组织、分层审核和快速上报场景的商业化产品，围绕“统一入口、AI生成、人工审核、统计考核、运营配置”五个环节，形成从事件发现到结果归档的完整闭环。', after=2)
add_body(doc, '当前系统已形成统一入口门户，以及 H5 报送审核端、V8 客户管理端、MT 应用管理端三大业务终端，可按客户组织结构、流程规范和部署环境灵活组合，满足“前端快速采集、后台统一配置、全程可追溯”的商业需求。', after=2)
add_body(doc, '系统的商业化定位主要体现在三个方面：', bold=True, size=11, after=3)
for t in [
    '面向组织级客户提供统一入口和多端协同工作台，降低多系统切换和培训成本。',
    '面向审核和运营场景提供 AI 辅助生成、同类事件识别、首发追溯、重复报送识别等能力。',
    '面向平台管理场景提供租户、权限、模板、规则、流程、统计与日志的配置化管理能力。',
]:
    add_body(doc, t, bullet=True, indent=0.3)

add_heading(doc, '二、行业现状', 1)
add_body(doc, '在传统速报、审核与统计类系统中，普遍存在入口分散、报送格式不统一、人工录入耗时、重复报送难识别、审核链路不透明、权限配置粗放等问题。对于多层级组织而言，报送数据往往分布在不同端、不同角色和不同数据口径中，导致运营侧难以快速汇总，审核侧难以追踪源头，管理侧难以形成有效考核。', after=3)
for t in [
    '一线上报依赖人工整理，容易出现模板不一致、字段缺失和附件遗漏。',
    '审核侧需要同时处理事件真实性、同类事件识别、首发追溯和重复报送判断，单靠人工效率低。',
    '平台侧需要兼顾机构、人员、二维码、角色权限、审核规则、考核规则和消息通知等多类配置。',
    '如果缺少统一入口和统一数据底座，后续统计、考核和运营分析都难以沉淀成可复用资产。',
]:
    add_body(doc, t, bullet=True)
add_body(doc, '点点速报系统正是围绕上述痛点设计，通过 AI 能力与流程编排，将“采集、审核、统计、考核、归档、追溯”整合为一个可交付、可运营、可扩展的产品。')

add_heading(doc, '三、产品介绍', 1)
add_body(doc, '点点速报系统采用“门户层 - 应用层 - 业务层 - 能力层 - 数据层”的分层架构。门户层面向不同角色提供统一入口；应用层组织报送中心、审核中心、统计中心、考核中心、消息中心与系统管理；业务层承载 H5 报送审核端、V8 客户管理端和 MT 应用管理端；能力层提供 AI、多模态分析、文件服务、消息服务和权限服务；数据层负责业务数据、分析数据、日志数据、缓存与搜索。', after=3)

add_heading(doc, '3.1 Web客户端', 2)
add_body(doc, 'Web 客户端主要面向机构管理员、审核管理员和平台运营人员，是点点速报系统的桌面级管理中枢。当前产品中，V8 客户管理端与 MT 应用管理端构成了 Web 侧的核心管理能力。', after=2)
for t in [
    'V8 客户管理端偏向业务运营与组织管理，支持组织架构、人员账号、角色权限、二维码管理、业务配置和系统日志。',
    'MT 应用管理端偏向平台级管理，覆盖机构信息、机构套餐、激活与审核模板、审核层级流程、拒绝理由、统计指标等关键参数。',
    'Web 端适合承担“看板 + 配置 + 审核 + 统计 + 考核”的复合型工作台角色。',
]:
    add_body(doc, t, bullet=True)

add_heading(doc, '3.2 H5端', 2)
add_body(doc, 'H5 端主要面向一线上报员和移动巡检人员，强调“随时可上报、随手可采集、随地可审核”的轻量化体验。其核心价值不是替代桌面端，而是在移动场景中尽可能缩短事件采集到进入审核的时间。', after=2)
for t in [
    '支持图文、语音、链接等多模态内容提交，满足现场快速采集需要。',
    '支持 AI 自动生成速报草稿，减少重复录入和格式整理工作。',
    '支持模板上报、报送草稿、报送记录、审核提醒和结果通知，形成完整的移动闭环。',
    '支持个人中心、基础信息认证、数据权限和统计查看，便于个人任务跟踪。',
]:
    add_body(doc, t, bullet=True)

add_heading(doc, '3.3 人工服务端', 2)
add_body(doc, '人工服务端主要承载审核、复核、统计、考核和运营管理等人工决策环节，是连接 AI 自动识别与最终业务落地的重要枢纽。商业版产品并不追求“全自动替代人工”，而是通过“AI 先筛、人工确认、流程留痕”提高整体效率与准确性。', after=2)
for t in [
    '支持审核中心、统计中心、考核中心和系统管理等业务模块。',
    '支持同类事件识别、首发识别追溯、重复报送识别与审核记录留存。',
    '支持机构、人员、角色、规则与流程的分层配置，便于不同客户按组织结构定制。',
    '支持消息中心、系统公告、审核提醒和结果通知，强化人工处置效率。',
]:
    add_body(doc, t, bullet=True)

add_heading(doc, '四、核心功能（市场商务方向）', 1)
add_heading(doc, '4.1 统一入口与多端协同工作台', 2)
add_body(doc, '点点速报系统的入口层不再是多个分散地址，而是一个统一工作台。用户进入门户后，可按角色和权限看到对应端入口，并通过新标签页或统一导航快速切换到 H5、V8 客户管理端、MT 应用管理端等业务终端。', after=2)
for t in [
    '统一品牌与统一视觉风格，降低客户对“多个系统”的割裂感。',
    '统一登录入口，减少重复认证和重复培训。',
    '按角色分发可见模块，避免无关功能干扰，提高使用效率。',
    '支持快速入口、消息提醒和常用功能直达，适合高频操作场景。',
]:
    add_body(doc, t, bullet=True)

add_heading(doc, '4.2 AI驱动的速报生成与审核闭环', 2)
add_body(doc, '系统将 AI 能力深度嵌入报送与审核流程，在“发现事件—提交事件—生成草稿—人工修订—提交审核—审核通过/驳回—归档统计”的链路中，帮助客户降低人工成本并提升审查效率。', after=2)
for t in [
    '多模态输入：支持图片、视频、音频、链接、文字等内容的识别与结构化提取。',
    'AI 草稿生成：根据事件分析结果自动形成速报文稿，显著减少首稿编写时间。',
    '智能识别能力：包含同类型事件识别、首发识别追溯、重复报送识别与不良信息库联动。',
    '审核辅助能力：对审核意见、审核记录、驳回原因、结果通知进行全过程管理。',
]:
    add_body(doc, t, bullet=True)

add_heading(doc, '4.3 多级组织、租户与权限运营', 2)
add_body(doc, '商业版点点速报系统特别强调多级组织和租户运营能力，能够覆盖平台、机构、子机构与用户的多层管理结构。平台侧负责统一策略和全局参数，机构侧负责业务承接和人员管理，子机构侧负责报送执行和本地运营。', after=2)
for t in [
    '支持机构信息、机构分类、人员信息、角色权限、二维码管理和组织架构管理。',
    '支持模板配置、激活模板配置、审核打分规则、拒绝理由、考核规则、统计指标和登录验证方式配置。',
    '支持机构套餐、增值业务开关、业务功能开关和平台权限配置，适合商业化售卖。',
    '支持系统日志、运营监控与权限中心，便于平台审计和客户运维。',
]:
    add_body(doc, t, bullet=True)

add_heading(doc, '4.4 业务流程与状态流转', 2)
add_body(doc, '从业务流程看，系统建立了清晰的状态流转模型：草稿 → 待提交 → 审核中 → 审核通过 / 被驳回 → 已归档。每一步都保留操作记录与结果留痕，便于后续统计、追责和分析。', after=2)
for t in [
    '上报员发现事件后，可先提交多模态素材，系统自动识别并生成报送草稿。',
    '人工修改补充后提交审核，进入基层审核和上级机构审核环节。',
    '审核时可结合 AI 同类型事件识别、首发识别追溯和重复报送识别进行判断。',
    '审核通过后进入归档、汇总统计、数据分析和过程留痕阶段，形成闭环。',
    '若审核被驳回，业务可回到修订环节，保证报送质量和责任清晰。',
]:
    add_body(doc, t, bullet=True)
add_body(doc, '这一套流程既适合“快速响应”场景，也适合“层级审核”场景，是点点速报系统区别于简单报送工具的核心竞争力。')

add_heading(doc, '五、适用场景', 1)
add_body(doc, '点点速报系统适用于需要多级组织协同、移动快速上报、人工审核留痕和指标考核管理的业务场景。', after=2)
for t in [
    '集团型企业与多分子机构管理场景：需要统一平台管理多个机构、多个角色和多个业务参数。',
    '现场巡检与移动报送场景：需要通过 H5 快速采集图文/语音/链接信息并迅速进入审核。',
    '风险审核与合规管理场景：需要对事件进行同类识别、首发追溯、重复报送判断与审核留痕。',
    '运营考核与统计分析场景：需要按机构、人员、区域、流程阶段形成统计、考核和导出能力。',
    '商业化平台交付场景：需要按客户套餐、功能开关和权限配置进行项目化或产品化交付。',
]:
    add_body(doc, t, bullet=True)

add_heading(doc, '六、部署与服务模式', 1)
add_body(doc, '商业版点点速报系统支持多种部署形态，可根据客户的安全要求、网络环境和运维能力进行灵活适配。', after=2)
add_kv_table(doc, [
    ('私有化部署', '系统部署在客户自有服务器或专属云环境中，数据和权限完全可控，适合对安全与合规要求较高的客户。'),
    ('专属云部署', '系统运行在独立租户环境内，兼顾交付效率与隔离性，适合需要快速上线的商业客户。'),
    ('混合部署', '核心数据和关键服务本地化部署，非敏感能力通过云端扩展，适合已有云化基础设施的客户。'),
    ('服务模式', '支持项目实施、流程梳理、权限配置、培训交付、持续运维和二次扩展。'),
])
add_body(doc, '在接入层面，系统支持浏览器访问、H5 移动端访问以及外部系统 API 对接，能够与指令流转、舆情分析、截图、短信、微信公众号、权限配置等外部系统协同工作。', after=2)
add_body(doc, '在商业服务层面，可按客户需求提供标准版、增强版和定制版方案，便于按模块售卖、按端售卖或按组织层级售卖。')

add_heading(doc, '七、产品优势', 1)
for t in [
    '入口统一：把分散的端统一到同一个门户，降低用户学习成本。',
    'AI 赋能：让系统从“报送工具”升级为“智能报送与审核平台”。',
    '流程闭环：从提交、审核、归档到统计分析，形成完整闭环。',
    '配置驱动：模板、规则、流程、权限和二维码都可配置，适配不同客户。',
    '多级组织能力强：支持平台、机构、子机构多层治理。',
    '商业化友好：支持套餐、功能开关、权限隔离和增值能力扩展。',
    '留痕可追溯：所有关键动作均保留记录，方便审计与复盘。',
]:
    add_body(doc, t, bullet=True)
add_body(doc, '对于客户而言，这意味着更短的上线周期、更稳定的流程管理和更清晰的数据资产沉淀。')

add_heading(doc, '八、客户价值', 1)
add_body(doc, '点点速报系统不仅是一套业务系统，更是一个帮助客户提升组织协同效率的数字化平台。其核心客户价值主要体现在以下几个方面：', after=2)
for t in [
    '提升报送效率：减少重复录入和格式整理，缩短从事件发现到提交审核的时间。',
    '提升审核效率：借助 AI 识别和规则配置，降低人工筛选压力。',
    '提升管理透明度：机构、人员、权限、流程和日志统一可视。',
    '提升数据可用性：形成结构化、可统计、可考核的数据资产。',
    '提升运营能力：支持按机构、按人员、按区域和按流程阶段进行精细化运营。',
    '提升商业拓展能力：通过模块化、可配置和可扩展的设计，支撑后续增值服务。',
]:
    add_body(doc, t, bullet=True)
add_body(doc, '从客户视角看，系统最终带来的不是单点功能升级，而是“报送更快、审核更准、统计更全、考核更清晰、运营更可控”的综合价值。')

add_heading(doc, '九、结语', 1)
add_body(doc, '点点速报系统商业版以统一入口门户为起点，以 H5 报送审核端、V8 客户管理端和 MT 应用管理端为核心载体，以 AI 能力、流程引擎和多级组织管理为支撑，构建了一套真正面向商业交付的速报平台。', after=2)
add_body(doc, '后续产品可在当前基础上继续扩展更多行业能力，例如更丰富的模型识别、更细粒度的流程编排、更灵活的计费体系和更完整的运营分析能力，使点点速报系统从“可用产品”持续演进为“可规模化复制的行业解决方案”。', after=2)
add_body(doc, '如果客户希望先快速落地，本产品也可以以“统一入口 + 多端独立部署”的方式交付；如果客户希望长期统一运营，则可进一步演进为单门户、多子系统和统一权限的完整产品体系。')

# Appendices
add_heading(doc, '附录A 产品总体架构图', 1)
add_body(doc, '该图展示了点点速报系统从外接系统、门户层、应用层、业务层、能力层到数据层的完整架构关系，可作为商业版方案讲解和方案汇报的总览页。', after=4)
if img1.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=2, line=1.0)
    p.add_run().add_picture(str(img1), width=Inches(6.6))
else:
    add_body(doc, f'未找到图片：{img1}', color='C00000')

add_page = doc.add_paragraph()
add_page.add_run().add_break(WD_BREAK.PAGE)

# Instead of add_page_break via paragraph to keep control
# Build next appendix on fresh page

add_heading(doc, '附录B 业务流程与状态流转图', 1)
add_body(doc, '该图展示了“发现事件—AI识别—草稿生成—人工修改—提交审核—审核流转—归档统计”的主流程，以及租户管理、机构管理与上报审核的状态链路。', after=4)
if img2.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=2, line=1.0)
    p.add_run().add_picture(str(img2), width=Inches(5.2))
else:
    add_body(doc, f'未找到图片：{img2}', color='C00000')

p = doc.add_paragraph()
p.add_run().add_break(WD_BREAK.PAGE)

add_heading(doc, '附录C 功能架构图', 1)
add_body(doc, '该图展示了点点速报系统的功能树总览，覆盖 H5 报送、V8 客户管理、MT 应用管理、系统配置、AI 能力以及相关消息与日志服务。', after=4)
if img3.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=2, line=1.0)
    p.add_run().add_picture(str(img3), width=Inches(6.6))
else:
    add_body(doc, f'未找到图片：{img3}', color='C00000')

# Final polish: set spacing for all paragraphs without explicit format a bit tighter
for para in doc.paragraphs:
    if para.style.name in ('Normal', 'List Bullet') and para.paragraph_format.space_after is None:
        set_paragraph_spacing(para, before=0, after=3, line=1.25)

# Core properties
cp = doc.core_properties
cp.title = '点点速报系统商业版产品白皮书'
cp.subject = '商业版产品白皮书'
cp.author = 'Codex'
cp.keywords = '点点速报系统, 产品白皮书, 商业版, 速报, H5, V8, MT'
cp.comments = 'Generated from product architecture and workflow references.'

# Save
# Remove accidental blank page-break-only paragraphs if possible not needed.
doc.save(str(DOCX))
print(str(DOCX))
